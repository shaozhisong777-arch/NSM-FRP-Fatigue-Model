from pathlib import Path
import math,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Cm,Pt as DPt
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
r100=15.72*(1-.1985*math.log10(100));r120=15.72*(1-.1985*math.log10(120))
slides=[
('汇报主线：疲劳怎样进入这套模型',
 ['反复荷载下，先求峰值与谷值的平衡响应。','局部滑移达到条件后，节点开始累计疲劳龄期。','龄期增加使峰值黏结应力降低，再重新求整体平衡。'],
 '局部滑移 → 活跃节点 → 局部龄期 → 黏结退化 → 响应重分配',
 '今天主要介绍疲劳计算过程。模型在固定幅值荷载下，采用循环块的准静态峰谷平衡计算。它用局部界面滑移判断疲劳是否激活，再通过节点黏结峰值退化，把循环历史传递到后续受力响应中。循环数增加并不意味着所有节点按同一程度退化。',
 'fatigue_calculate：主控制；active_nodes、bond_strength：局部疲劳规则'),
('先区分三个循环量和一份历史',
 ['N：整体已经接受的循环数。','Ni：某根筋、某一侧、某一节点的局部循环数。','ΔN：本次试算推进的循环块长度。','memory：最大滑移、局部龄期、激活和失效等已接受历史。'],
 '整体循环数 N ≠ 所有节点共享的局部循环数 Ni',
 '这里有三个容易混淆的量。N是整体计算已经接受的循环数，Ni是各节点的局部龄期，ΔN是本步打算推进多少次。一个节点只有处于活跃状态时才增加Ni。memory保存不可逆历史，也是计算重试时回到的基准。虽然候选数据可能反复计算，已接受的历史不会随意被覆盖。',
 'fatigue_calculate：create_memory；update_memory'),
('第1步：建立初始峰谷响应',
 ['先检查参数、几何和交点，初始化历史为零。','求初始峰值平衡，保存该状态的最大滑移历史。','继承峰值历史，再求初始谷值平衡。','峰谷均成功，才记录N=0并进入疲劳循环。'],
 'Vc + Vs + Vf = Qmax/2；   Vc + Vs + Vf = Qmin/2',
 '疲劳更新之前，必须先知道这个荷载范围下的初始峰谷响应。程序先求峰值平衡，再用峰值的最大滑移历史求谷值。两次计算的目的，是取得每个节点的峰谷滑移差，而不是只看峰值滑移。N等于零表示尚未提交疲劳循环更新，不表示没有外荷载。',
 'fatigue_calculate第18–54行；state → trial'),
('第2步：判断活跃节点并提出局部龄期更新',
 ['资格：smax > 0.75δm，或者该节点曾经激活。','同时要求：max(0, speak − slow) > cyclicSlipTolerance。','活跃节点Ni增加ΔN；其他节点的Ni保持原值。','先复制历史，生成候选；此时还不能提交。'],
 'Ni,proposed = Ni,accepted + ΔN（仅活跃节点）',
 '当前程序的活跃判断有两道门槛。节点需要达到滑移阈值或已经激活，同时本步仍然有足够的峰谷滑移差。因此，以前激活过的节点在没有循环滑移时，也不会继续计数。确定活跃节点后，update_memory只在这些节点增加局部龄期。0.75阈值有参考论文依据，历史保持与滑移差容差是代码实现规则。',
 'active_nodes；fatigue_calculate第58行及第232–246行'),
('第3步：把局部龄期转成黏结退化',
 ['杆材kf=0.1965；条材kf=0.1985。','采用Ni而不是直接对所有节点使用整体N。','局部龄期为0时，用max(Ni,1)避免对数无定义。','峰值取非负值；该规律针对界面，不是FRP本体疲劳强度。'],
 'τm,i = max{0, τm,0[1 − kf log₁₀max(Ni,1)]}',
 f'接下来，程序把局部龄期转换成节点的峰值黏结应力。这一步的公式在bond_strength中，实际通过state、trial被调用，不是在update_memory里直接完成。以当前条材初始峰值15.72MPa为例，同一活跃节点由100次推进到120次后，峰值从约{r100:.3f}MPa降到{r120:.3f}MPa。这只是公式代入示例，不是一组新增的试验结果或整梁模拟结果。',
 'bond_strength；Chen, Wang & Cheng (2019), IJ Fatigue, Eq. (2), DOI:10.1016/j.ijfatigue.2019.04.039'),
('第4步：退化后重新平衡，决定接受还是重试',
 ['候选Ni → 新τm → 新黏结关系 → 锚固受力 → 新峰值平衡。','峰值失败或q增长过快，且ΔN>1时，循环块减半后重算。','峰值成功，再用其历史求候选谷值；两者成功才提交。','接受后：N增加ΔN，替换memory，并记录历史与快照。'],
 'proposed = update_memory(...) → state(...proposed...) → accept / retry',
 '疲劳更新之后，不能直接假定结构仍然平衡。程序需要用退化后的黏结关系重算峰值响应。如果失败或位移跳得过大，就在循环块大于一次的前提下减半。每次重试都从同一份已接受历史出发，不把错误候选累积进去。峰值通过后，还要检查谷值。当前代码只对上述峰值情况减半，谷值失败则返回诊断。两状态都成功，才真正提交这一步。',
 'fatigue_calculate第106–155行；state → trial → bond_strength / anchorage'),
('第5步：循环块控制与停止条件',
 ['ΔN不跨越Nmax、下一快照位置或设定cycleStep。','活跃节点集合由已接受峰谷状态确定，块内不逐周期重新判断。','无活跃节点且其他疲劳规则均关闭：记录至Nmax。','数值失败、模型限制与runout分别记录。'],
 '达到 Nmax 是 runout；单根筋失效后仍需检查整体平衡',
 '循环块是数值近似，作用是提高效率。程序不逐周期连续解析完整波形，也不是完整滞回时程积分。在一个块内，活跃节点按块起点的峰谷响应确定，因此需要检查循环步长敏感性。达到Nmax只说明算到了设定上限。如果最小块仍找不到指定域内的平衡，或局部求解不收敛，应报告相应原因，不能全部叫物理脱黏破坏。',
 'fatigue_calculate第72–100行、第117–145行及最终状态整理'),
('明天汇报的结论与当前范围',
 ['已经实现：局部激活、局部计数、界面黏结退化和峰谷重平衡。','默认关闭：箍筋、混凝土S–N/Miner疲劳更新。','尚未实现：FRP本体疲劳断裂，以及完整混凝土疲劳破坏机制。','后续验证：局部退化区网格收敛、循环步长、抗剪疲劳试验对比。'],
 '建议表述：基于局部界面疲劳退化的抗剪平衡迭代计算框架',
 '本次汇报着重说明计算框架已经怎样运行，而不是宣称所有破坏机制已经完整建立。当前默认主要研究局部界面疲劳退化引起的响应变化。混凝土提供经验抗剪贡献，FRP有固定抗拉上限，但它们不等于完整的混凝土和FRP本体疲劳模型。此外，前面的审查提示疲劳增量对网格较敏感，因此下一阶段需要继续收敛研究和试验验证。',
 '当前上传fatigue_calculate与仓库计算实现一致；本次仅制作汇报材料，未修改源码'),
]
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def txt(s,x,y,w,h,t,size=22,color='243746',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 p=tf.paragraphs[0];p.text=t;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
 for run in p.runs:
  ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
 return sh
for i,(title,bullets,key,note,where) in enumerate(slides,1):
 s=prs.slides.add_slide(prs.slide_layouts[6]);txt(s,.6,.35,12,.65,title,28,'173A5E',True)
 for k,b in enumerate(bullets):txt(s,.75,1.5+k*.87,11.8,.72,'• '+b,23)
 band=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.65),Inches(5.3),Inches(12.0),Inches(1.05));band.fill.solid();band.fill.fore_color.rgb=RGBColor.from_string('EAF2FF');band.line.fill.background();txt(s,.95,5.55,11.4,.63,key,22,'173A5E',True)
 txt(s,.65,6.66,11.8,.5,where,12,'64758B');txt(s,12.1,7.08,.6,.2,f'{i}/8',11,'64758B')
 s.notes_slide.notes_text_frame.text=note
prs.save(ROOT/'疲劳计算过程_汇报重点_CN.pptx')
md=['# 疲劳计算过程：明天汇报提纲与讲稿','本材料对应本次上传的fatigue_calculate.m。核心代码与当前版本一致。建议讲述时间约5—7分钟，取决于公式解释程度；PPT每页备注也已写入对应讲稿。','']
for i,(title,bullets,key,note,where) in enumerate(slides,1):
 md += [f'## 第{i}页：{title}','',note,'',f'核心句：{key}',f'代码定位：{where}','']
md+=['## 公式代入示例（不是新增试验或整梁模拟结果）','',f'条材初始峰值15.72 MPa、kf=0.1985：Ni=100时τm={r100:.6f} MPa；Ni=120时τm={r120:.6f} MPa。只有本步活跃的节点才增加局部龄期。','','## 一分钟总结','', '本程序先求当前荷载范围下的峰值和谷值响应，再通过滑移阈值与峰谷滑移差识别疲劳活跃节点。只给这些节点累计局部循环数，并使用局部疲劳退化律降低峰值黏结应力。随后重新求峰谷平衡；必要时减小循环块重试。两状态都成功才更新整体循环数与已接受历史。这是一套基于局部界面退化的准静态抗剪疲劳迭代框架，后续仍需网格收敛与试验验证。']
text='\n'.join(md);(ROOT/'疲劳计算过程_中文讲稿.md').write_text(text)
d=Document();sec=d.sections[0];sec.left_margin=sec.right_margin=Cm(2.4);sec.top_margin=sec.bottom_margin=Cm(2.0)
for name in ['Normal','Title','Heading 1']:
 st=d.styles[name];st.font.name='Calibri';st.font.size=DPt(12 if name=='Normal' else 16);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
d.styles['Normal'].paragraph_format.line_spacing=1.3
for block in text.split('\n\n'):
 if block.startswith('# '):d.add_heading(block[2:],0)
 elif block.startswith('## '):d.add_heading(block[3:],1)
 else:d.add_paragraph(block)
d.save(ROOT/'疲劳计算过程_中文讲稿.docx');(ROOT/'slide_notes.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2))
print('Illustrative strip peak stress:',r100,'→',r120,'MPa; 8 slides with speaker notes and Word script.')
