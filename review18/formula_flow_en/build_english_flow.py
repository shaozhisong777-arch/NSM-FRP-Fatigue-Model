from pathlib import Path
import json,html,math
from xml.etree.ElementTree import Element,SubElement,ElementTree
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
r=Path(__file__).resolve().parent
BLUE='E7F0FF';GREEN='E5F3E9';ORANGE='FFE9CD';RED='FFE2DE';GRAY='EEF0F3'

def N(id,text,y,h=100,kind='box',color=GRAY,x=430,w=570):return dict(id=id,text=text,x=x,y=y,w=w,h=h,kind=kind,color=color)
def E(a,b,label='',ap='b',bp='t',via=None):return dict(a=a,b=b,label=label,ap=ap,bp=bp,via=via or [])

# F=2019 fatigue paper; S=2019 single-crack paper; E=2024 ECC paper.
diagrams=json.loads((r/'diagram_data.json').read_text())
footers=['F: 2019 flexural-fatigue modeling paper; S: 2019 single-crack shear paper; E: 2024 ECC-FRP paper. Equation numbers need a paper prefix.', 'F: DOI 10.1016/j.ijfatigue.2019.04.039; S: DOI suffix 0000953; E: CCENG-4395. Orange boxes include model assumptions.', 'R22: DOI suffix 0001251, intersecting reinforcement and shear-force decomposition. This update translates the diagrams only.']
mx=Element('mxfile',host='app.diagrams.net',type='device')
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(28)
for index,d in enumerate(diagrams):
 nodes={n['id']:n for n in d['nodes']}
 def port(n,p):
  return {'t':(n['x']+n['w']/2,n['y']),'b':(n['x']+n['w']/2,n['y']+n['h']),'l':(n['x'],n['y']+n['h']/2),'r':(n['x']+n['w'],n['y']+n['h']/2)}[p]
 md=SubElement(mx,'diagram',name=d['name'],id=f'calc{index}')
 model=SubElement(md,'mxGraphModel',grid='1',gridSize='10',page='1',pageWidth=str(d['width']),pageHeight=str(d['height']))
 root=SubElement(model,'root');SubElement(root,'mxCell',id='0');SubElement(root,'mxCell',id='1',parent='0')
 def celltext(id,s,x,y,w,h,font=22):
  c=SubElement(root,'mxCell',id=id,value=s,style=f'text;html=0;align=center;verticalAlign=middle;whiteSpace=wrap;fontFamily=Arial;fontSize={font};',vertex='1',parent='1');SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
 celltext('title',d['name'],0,5,d['width'],45,28)
 celltext('footer',footers[index],25,d['height']-80,d['width']-50,65,18)
 for n in d['nodes']:
  shape='rhombus' if n['kind']=='diamond' else 'rectangle'
  style=f'shape={shape};rounded=1;whiteSpace=wrap;html=0;fillColor=#{n["color"]};strokeColor=#6A7E90;fontFamily=Arial;fontSize={18 if n["w"]>400 else 16};'
  c=SubElement(root,'mxCell',id=n['id'],value=n['text'],style=style,vertex='1',parent='1');SubElement(c,'mxGeometry',x=str(n['x']),y=str(n['y']),width=str(n['w']),height=str(n['h']),attrib={'as':'geometry'})
 for i,e in enumerate(d['edges']):
  style='edgeStyle=orthogonalEdgeStyle;rounded=0;endArrow=block;endFill=1;strokeColor=#1685B6;fontColor=#B8302E;fontSize=19;'
  pp={'t':(.5,0),'b':(.5,1),'l':(0,.5),'r':(1,.5)};a=pp[e['ap']];b=pp[e['bp']];style+=f'exitX={a[0]};exitY={a[1]};entryX={b[0]};entryY={b[1]};'
  c=SubElement(root,'mxCell',id='e'+str(i),value=e['label'],style=style,edge='1',parent='1',source=e['a'],target=e['b']);geo=SubElement(c,'mxGeometry',relative='1',attrib={'as':'geometry'})
  if e['via']:
   arr=SubElement(geo,'Array',attrib={'as':'points'})
   for x,y in e['via']:SubElement(arr,'mxPoint',x=str(x),y=str(y))
 # Preview SVG uses identical editable vector primitives and labels.
 svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{d["width"]}" height="{d["height"]}" viewBox="0 0 {d["width"]} {d["height"]}">','<defs><marker id="a" markerWidth="9" markerHeight="9" refX="8" refY="4.5" orient="auto"><path d="M0,0 L9,4.5 L0,9Z" fill="#1685B6"/></marker></defs>','<rect width="100%" height="100%" fill="white"/>']
 def svgtext(s,cx,cy,fs=21,color='#243746'):
  lines=s.split('\n');out=[]
  for i,line in enumerate(lines):out.append(f'<text x="{cx}" y="{cy+(i-(len(lines)-1)/2)*fs*1.32}" font-family="Arial,DejaVu Sans,sans-serif" font-size="{fs}" text-anchor="middle" dominant-baseline="middle" fill="{color}">{html.escape(line)}</text>')
  return out
 svg+=svgtext(d['name'],780,30,28)
 for e in d['edges']:
  pts=[port(nodes[e['a']],e['ap'])]+e['via']+[port(nodes[e['b']],e['bp'])]
  svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in pts)+'" fill="none" stroke="#1685B6" stroke-width="3" marker-end="url(#a)"/>')
  tip=pts[-1];prev=pts[-2];dx=tip[0]-prev[0];dy=tip[1]-prev[1];length=math.hypot(dx,dy)
  if length:
   ux,uy=dx/length,dy/length;bx,by=tip[0]-18*ux,tip[1]-18*uy
   svg.append(f'<polygon points="{tip[0]},{tip[1]} {bx-7*uy},{by+7*ux} {bx+7*uy},{by-7*ux}" fill="#1685B6"/>')
  if e['label']:
   p0,p1=pts[0],pts[1];lx,ly=e.get('lp',[(p0[0]+p1[0])/2+18,(p0[1]+p1[1])/2-12]);svg+=svgtext(e['label'],lx,ly,19,'#B8302E')
 for n in d['nodes']:
  x,y,w,h=n['x'],n['y'],n['w'],n['h']
  if n['kind']=='diamond':svg.append(f'<polygon points="{x+w/2},{y} {x+w},{y+h/2} {x+w/2},{y+h} {x},{y+h/2}" fill="#{n["color"]}" stroke="#6A7E90"/>')
  else:svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="14" fill="#{n["color"]}" stroke="#6A7E90"/>')
  svg+=svgtext(n['text'],x+w/2,y+h/2,18 if n['w']>400 else 16)
 svg+=svgtext(footers[index],780,d['height']-45,17);svg.append('</svg>');(r/(d['name']+'.svg')).write_text('\n'.join(svg))
 # Standard PowerPoint nodes and native straight-segment connectors; no custom paths.
 slide=prs.slides.add_slide(prs.slide_layouts[6]);scale=min(15/d['width'],26.4/d['height']);ox=(16-d['width']*scale)/2;oy=(28-d['height']*scale)/2
 def pos(x,y):return Inches(ox+x*scale),Inches(oy+y*scale)
 def ppttext(sh,s,fs=21):
  tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.margin_top=tf.margin_bottom=0;tf.margin_left=tf.margin_right=0
  for i,t in enumerate(s.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t;p.alignment=PP_ALIGN.CENTER;p.font.name='Arial';p.font.size=Pt(fs*scale*72);p.space_after=Pt(1);p.font.color.rgb=RGBColor.from_string('243746')
 shapes={}
 for n in d['nodes']:
  kind=MSO_SHAPE.DIAMOND if n['kind']=='diamond' else MSO_SHAPE.ROUNDED_RECTANGLE
  sh=slide.shapes.add_shape(kind,*pos(n['x'],n['y']),Inches(n['w']*scale),Inches(n['h']*scale));sh.name=n['id'];sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(n['color']);sh.line.color.rgb=RGBColor.from_string('6A7E90');ppttext(sh,n['text'],18 if n['w']>400 else 16);shapes[n['id']]=sh
 for e in d['edges']:
  pts=[port(nodes[e['a']],e['ap'])]+e['via']+[port(nodes[e['b']],e['bp'])]
  for seg,(p0,p1) in enumerate(zip(pts,pts[1:])):
   line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*pos(*p0),*pos(*p1));line.name=e['a']+' → '+e['b']+f' / {seg+1}';line.line.color.rgb=RGBColor.from_string('1685B6');line.line.width=Pt(1.4)
   sites={'t':0,'l':1,'b':2,'r':3}
   if seg==0:line.begin_connect(shapes[e['a']],sites[e['ap']])
   if seg==len(pts)-2:
    line.end_connect(shapes[e['b']],sites[e['bp']]);tip=OxmlElement('a:tailEnd');tip.set('type','triangle');line._element.spPr.get_or_add_ln().append(tip)
   tree=slide.shapes._spTree;tree.remove(line._element);tree.insert(2,line._element)
  if e['label']:
   p0,p1=pts[0],pts[1];x,y=e.get('lp',[(p0[0]+p1[0])/2+18,(p0[1]+p1[1])/2-20])
   tb=slide.shapes.add_textbox(*pos(x-100,y-15),Inches(200*scale),Inches(35*scale));ppttext(tb,e['label'],18)
 for txt,x,y,w,h,fs in [(d['name'],0,5,1560,45,28),(footers[index],10,d['height']-80,1540,65,17)]:
  tb=slide.shapes.add_textbox(*pos(x,y),Inches(w*scale),Inches(h*scale));ppttext(tb,txt,fs)
ElementTree(mx).write(r/'formula_mapped_shear_fatigue_editable_EN.drawio',encoding='utf-8',xml_declaration=True)
prs.save(r/'formula_mapped_shear_fatigue_editable_EN.pptx')
print('Built three English formula-mapped flowcharts.')
