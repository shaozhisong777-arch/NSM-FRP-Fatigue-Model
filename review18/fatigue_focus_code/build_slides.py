from pathlib import Path
import textwrap,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
ROOT=Path(__file__).resolve().parent
REPO=ROOT.parent.parent
sources={'fatigue_calculate.m':(ROOT/'fatigue_calculate_uploaded.m.txt').read_text().splitlines()}
for file in ['active_nodes.m','bond_strength.m','parameters.m']:
 sources[file]=(REPO/file).read_text().splitlines()
# All displayed MATLAB lines are selected from actual files, with original line numbers.
slides=[
('主线：局部疲劳怎样推动整体响应','fatigue_calculate.m',[(58,58),(101,102),(147,151)],
['先识别活跃节点，提出局部疲劳历史。','用候选历史重算响应，成功后才提交。','接受结果后，整体循环数才向前推进。'],
'判断 → 提出候选 → 重新平衡 → 接受',
'开场先用这几行代码说明整个结构。active_nodes给出节点标志，update_memory提出候选历史，state计算候选响应。后面几行在峰谷都成功后才执行，体现先验证、再更新。右侧选段之间的省略符表示中间代码未展示。'),
('N、Ni、ΔN与memory分别是什么','fatigue_calculate.m',[(191,197),(201,203)],
['N：整体已经接受的循环数。','Ni：保存在memory.age中的节点局部龄期。','ΔN对应cycleBlock；memory保留不可逆历史。'],
'每根筋、每一侧、每个节点都有自己的Ni',
'右侧create_memory局部函数建立了需要保存的历史。age按筋与侧分成cell，每个cell内是节点列向量。bondMax保存最大滑移，activated保存激活记录，failedFRP保存失效标志。Ni不是所有节点都相同的整体循环数。'),
('初始状态：先峰值，再谷值','fatigue_calculate.m',[(29,30),(44,46),(52,54)],
['峰值目标剪力为Qmax/2。','峰值的最大滑移历史传给谷值计算。','初始峰谷均成功后记录N=0。'],
'先确定当前荷载范围下的局部滑移',
'在进入疲劳循环前，先求峰值，再用峰值历史求谷值。后续激活判定需要峰谷滑移差。展示代码省略了两处失败返回分支，实际只有求解成功才会继续。N等于零表示尚未提交疲劳更新，并非没有施加外荷载。'),
('节点激活：本步哪些位置需要计数','active_nodes.m',[(34,35),(38,41),(44,48)],
['超过滑移阈值，或者以前已经激活。','本步峰谷滑移差还必须超过容差。','这里仅判断标志，暂不增加循环数。'],
'资格满足 AND 当前仍有循环滑移',
'本页为配套active_nodes.m的真实代码，不是这次粘贴主文件中的局部函数。activationSlip在该文件前面定义为activationFraction乘deltaM，默认比例0.75。两个条件必须同时成立。之前激活过但本步没有足够滑移差的节点，也不继续累计龄期。'),
('选循环块，并更新候选局部龄期','fatigue_calculate.m',[(95,96),(232,234),(240,243)],
['循环块不跨越Nmax或下一快照位置。','先复制memory，保留已接受的历史。','仅活跃节点的Ni增加cycleBlock。'],
'候选更新不等于已经接受',
'前两行控制本步推进多少次。proposed等于memory是在函数内复制已接受状态。下面的if位于筋、侧和节点三层循环内，只更新活跃节点。该片段没有把退化公式直接写进update_memory，而是先更新局部计数。'),
('黏结退化：把Ni转成峰值黏结应力','bond_strength.m',[(17,17),(20,24)],
['τm,0是初始峰值；kf按杆材或条材选取。','用max(Ni,1)避免零龄期的对数问题。','退化峰值非负，降低的是界面黏结强度。'],
'τm,i=max{0,τm,0[1−kf log₁₀max(Ni,1)]}',
'本页来自配套bond_strength.m。实际通过state、trial调用到这里，不在主控制器中直接降低参数。公式对应Chen、Wang和Cheng在2019年疲劳论文中的式2，杆材系数0.1965，条材0.1985。它是界面峰值退化，不是FRP本体抗拉强度的疲劳律。'),
('候选峰值：退化后重新求平衡','fatigue_calculate.m',[(101,109)],
['把候选历史proposed传给峰值求解。','state通过trial使用更新后的局部黏结峰值。','检查新峰值q是否增长过快。'],
'退化 → 锚固受力变化 → 重新寻找平衡q',
'第一行提出候选，第二行用候选历史求峰值平衡。之后只有求解成功才检查q增长量。这里q是裂缝竖向位移差，maxOpeningStep控制的是q增长尝试，不是直接限制法向裂缝宽度。求解得到的响应不会在这里立即提交。'),
('循环块减半：从原历史重新试算','fatigue_calculate.m',[(112,121)],
['峰值失败或q增长过快时触发重试。','前提是cycleBlock大于1。','减半后回到原已接受历史生成新候选。'],
'减半改善数值分辨率，不是材料损伤公式',
'这段控制是否再次进入retryBlock循环。只要循环块大于1，峰值失败或增长过快都会减半。再次执行update_memory时仍从原memory出发，不在之前失败候选上继续叠加。循环块达到1后，即使位移增长仍大也不再减半；是否能继续取决于峰值平衡是否成功。'),
('谷值通过后，才真正提交循环块','fatigue_calculate.m',[(138,144),(147,151)],
['候选峰值成功后，再计算候选谷值。','谷值失败直接返回诊断。','两者成功才更新N、peak、low和memory。'],
'峰值成功 AND 谷值成功 → 接受',
'nextMemory继承候选峰值的历史，再以Qmin/2求谷值。谷值失败则break，此时不会执行下面的提交。只有两状态都成功，N才增加，新的峰谷状态才替换旧状态。主循环的峰值减半机制并没有同时用于谷值失败。'),
('记录结果：历史与快照各有用途','fatigue_calculate.m',[(151,151),(158,163),(170,173)],
['history记录循环数、响应分量与残差等。','snapshots仅在指定循环数保存完整局部场。','out保留最后已接受的峰谷状态和历史。'],
'记录的是已接受结果，不是失败候选',
'这部分负责结果可追溯。history按每个接受块记录，而snapshots只在指定循环位置保留较完整的数据。结束后out返回最后的平衡状态和历史。这里的行数并不代表逐周期输出了连续时程。'),
('怎样停止：模型限制、数值失败与runout','fatigue_calculate.m',[(125,136),(167,169)],
['峰值求解失败时记录具体status。','指定模型或用户限制可记录失败循环区间。','达到Nmax记录runout，不当作破坏寿命。'],
'停止原因必须分别解释',
'代码会先检查失败属于哪一类。132至133行位于原文件的if modelLimitReached or strutLimitReached or momentLimitReached条件内，仅对相应限制记录失败循环及区间，不对所有数值失败都赋予物理寿命。右侧保留了分类条件，备注也提供原代码供核对。running在正常到达循环上限后改为runout_at_Nmax。'),
('默认范围：当前主要启用界面疲劳','parameters.m',[],
['箍筋、混凝土S–N参考循环数默认Inf。','相关可选更新因此不会执行。','FRP本体疲劳和完整混凝土破坏仍需补充。'],
'可以讲清计算框架，保留后续模型验证空间',
'本页给出当前parameters中两个Inf设置。主控制器在update_memory中用isfinite判断是否启用相应规则，默认都关闭。因此明天可重点汇报局部界面退化怎样推动峰谷响应变化，不应声称所有材料的疲劳破坏都已经实现。网格客观性与试验验证仍是后续重点。'),
]
# Select exact parameter lines by their field names, not fixed historical numbering.
param_lines=sources['parameters.m'];last=slides[-1];idx=[i+1 for i,t in enumerate(param_lines) if 'p.steelSNrefCycles =' in t or 'p.concreteSNrefCycles =' in t];slides[-1]=(last[0],last[1],[(i,i) for i in idx],*last[3:])
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
def text(s,x,y,w,h,string,size=22,color='263A4D',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 p=tf.paragraphs[0];p.text=string;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
 for run in p.runs:
  ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
 return sh
manifest=[]
for i,(title,file,ranges,bullets,key,note) in enumerate(slides,1):
 s=prs.slides.add_slide(prs.slide_layouts[6]);text(s,.6,.35,14.8,.7,title,29,'173A5E',True)
 text(s,.65,1.28,4.95,.35,'计算含义',17,'64758B',True)
 for k,b in enumerate(bullets):text(s,.65,1.94+k*1.2,4.95,1.05,'• '+b,23)
 box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.55),Inches(6.1),Inches(5.15),Inches(1.28));box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string('EAF2FF');box.line.fill.background();text(s,.8,6.36,4.65,.83,key,20,'173A5E',True)
 panel=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(6.05),Inches(1.4),Inches(9.35),Inches(6.43));panel.fill.solid();panel.fill.fore_color.rgb=RGBColor.from_string('182332');panel.line.fill.background()
 text(s,6.35,1.66,8.75,.4,file+'  |  原文件行号',18,'9EC9F5',True)
 raw=[];record=[]
 for j,(a,b) in enumerate(ranges):
  if j:raw.append((None,'…'))
  segment=sources[file][a-1:b];nonempty=[len(t)-len(t.lstrip()) for t in segment if t.strip()];indent=min(nonempty) if nonempty else 0
  for ln,t in enumerate(segment,a):
   code=t[indent:] if t.strip() else '';raw.append((ln,code));record.append((ln,t))
 lines=[]
 for ln,code in raw:
  prefix=f'{ln:>3}  ' if ln is not None else '     '
  if not code:lines.append((prefix,'dim'));continue
  wraps=textwrap.wrap(code,width=63,break_long_words=False,break_on_hyphens=False,replace_whitespace=False,drop_whitespace=False) or ['']
  for z,chunk in enumerate(wraps):
   kind='highlight' if any(term in chunk for term in ['proposed.age','proposed =','N = N +','tauM(node)','peakStress =','cycleBlock = floor','memory = nextMemory','out.status =']) else ('dim' if ln is None or chunk.lstrip().startswith('%') else 'code')
   lines.append(((prefix if z==0 else '     ↳ ')+chunk,kind))
 fs=min(15.0,5.08*72/(max(1,len(lines))*1.22));assert fs>=11.5,(i,fs,len(lines))
 codebox=s.shapes.add_textbox(Inches(6.35),Inches(2.33),Inches(8.76),Inches(5.12));tf=codebox.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for j,(line,kind) in enumerate(lines):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='Consolas';p.font.size=Pt(fs);p.line_spacing=Pt(fs*1.22);p.space_after=Pt(0);p.font.color.rgb=RGBColor.from_string({'code':'E4EBF3','dim':'8193A8','highlight':'FFD080'}[kind])
 text(s,.65,8.23,14,.35,'代码为真实选段；…表示省略，↳表示展示折行。源码和计算公式未修改。',13,'64758B');text(s,14.7,8.23,.7,.35,f'{i:02d}/12',13,'64758B')
 exact='\n'.join(f'{ln:>3}: {t}' for ln,t in record)
 s.notes_slide.notes_text_frame.text=note+'\n\n来源：'+file+'\n'+exact
 if i==11:s.notes_slide.notes_text_frame.text+='\n\n完整分类条件：\n'+'\n'.join(sources[file][124:136])
 manifest.append(dict(slide=i,title=title,file=file,ranges=ranges,font_size=fs,display_lines=len(lines)))
prs.save(ROOT/'疲劳计算过程_讲解与代码对照_CN.pptx')
(ROOT/'code_excerpt_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print('Built 12 slides with editable code and speaker notes. Minimum code font:',min(m['font_size'] for m in manifest))
