from pathlib import Path
import json,html,math
from xml.etree.ElementTree import Element,SubElement,ElementTree
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
r=Path(__file__).resolve().parent
BLUE='E7F0FF';GREEN='E5F3E9';ORANGE='FFE9CD';RED='FFE2DE';GRAY='EEF0F3'

def N(id,text,y,h=100,kind='box',color=GRAY,x=430,w=570):return dict(id=id,text=text,x=x,y=y,w=w,h=h,kind=kind,color=color)
def E(a,b,label='',ap='b',bp='t',via=None):return dict(a=a,b=b,label=label,ap=ap,bp=bp,via=via or [])

# F=2019 fatigue paper; S=2019 single-crack paper; E=2024 ECC paper.
main=[
N('start','开始：输入真实试件与峰谷总荷载 Qmax、Qmin',60,80,'round',ORANGE),
N('input','准备裂缝与筋的真实交点、两侧锚固长度\n初始化 N=0、局部龄期 Ni=0、最大滑移和失效记录\nparameters / prepare_parameters / fatigue_calculate',185,120,color=BLUE),
N('initial','依次求初始峰值和谷值抗剪平衡\nVc(q)+Vs(q)+Vf(q)=Qmax/2 或 Qmin/2\nstate → trial；内部计算见第2页',350,120,color=GREEN),
N('ok0','初始峰、谷求解成功？',510,100,'diamond',ORANGE),
N('limit','N < Nmax？',655,100,'diamond',ORANGE),
N('active','根据峰谷滑移判断本步活跃节点\n(smax>0.75δm 或曾激活)，且峰谷滑移差>容差\nactive_nodes｜F：阈值；历史保持与容差为代码规则',795,120,color=BLUE),
N('on','有活跃节点或启用了\n其他材料 S–N 退化？',960,110,'diamond',ORANGE),
N('step','选取循环块 ΔN，不跨越下一快照或 Nmax\n从已接受历史提出候选：Ni,new=Ni,old+ΔN\n仅活跃节点增加龄期；候选尚不提交',1110,120,color=BLUE),
N('proposal','更新节点峰值黏结应力 [F式2]\nτm,i=max{0,[1−kf log₁₀max(Ni,1)]τm,0}\nkf=0.1965（杆）或0.1985（条）｜bond_strength',1270,120,color=GREEN),
N('peak','将 τm,i 代入黏结本构，重新求峰值平衡\n由裂缝位移q → 滑移场 → 筋力 → 抗剪力\nstate → trial → anchorage / bond（见第2页）',1430,120,color=GREEN),
N('retry','峰值失败或 q 增长过快，\n且 ΔN > 1？',1590,110,'diamond',ORANGE),
N('peakok','候选峰值平衡成功？',1740,100,'diamond',ORANGE),
N('low','继承峰值最大滑移历史，求谷值平衡\n当前FRP黏结应力＝历史割线刚度×当前滑移\nstate → trial；谷值目标剪力 Qmin/2',1880,120,color=GREEN),
N('lowok','候选谷值平衡成功？',2040,100,'diamond',ORANGE),
N('accept','接受本循环块：N←N+ΔN，提交候选历史\n保存 q、w、v、Vc、Vs、Vf 与局部滑移/应变/失效\nfatigue_calculate / save_snapshot',2180,120,color=BLUE),
N('output','停止并输出最后有效状态\nexport_results；达到Nmax为runout，不等于疲劳寿命',2360,110,'round',ORANGE),
N('half','减小循环块\nΔN←max[1,floor(ΔN/2)]\n从同一份已接受历史重算',1255,140,color=BLUE,x=1120,w=330),
N('stop','记录停止原因\n局部数值失败 / 用户限制 /\n域内指定分支无平衡根\n不能统一写成“脱黏破坏”',2040,160,color=RED,x=1120,w=330),
N('runout','达到循环上限\n记录runout',655,100,color=ORANGE,x=50,w=300),
N('inactive','没有可演化退化项\n记录快照至Nmax\n无活跃退化runout',960,120,color=ORANGE,x=50,w=300)
]
edges=[E('start','input'),E('input','initial'),E('initial','ok0'),E('ok0','limit','是：保存N=0'),E('ok0','stop','否','r','r',[[1510,560],[1510,2120]]),E('limit','active','是'),E('limit','runout','否','l','r'),E('runout','output','','l','l',[[20,705],[20,2415]]),E('active','on'),E('on','step','是'),E('on','inactive','否','l','r'),E('inactive','output','','l','l',[[20,1020],[20,2415]]),E('step','proposal'),E('proposal','peak'),E('peak','retry'),E('retry','half','是','r','r',[[1475,1645],[1475,1325]]),E('half','proposal','','l','r',[[1060,1325],[1060,1330]]),E('retry','peakok','否'),E('peakok','low','是'),E('peakok','stop','否','r','l',[[1070,1790],[1070,2120]]),E('low','lowok'),E('lowok','accept','是'),E('lowok','stop','否','r','l'),E('accept','limit','下一循环块','l','l',[[375,2240],[375,705]]),E('stop','output','','b','r',[[1285,2415]])]
for edge in edges:
 if edge['a']=='accept':edge['lp']=[290,1700]
local=[
N('start','给定峰/谷目标剪力与已接受历史，选取试探 q\nstate：扫描上升符号变化及二分；不提交试探历史',65,105,'round',BLUE),
N('crack','① 裂缝运动（当前代码的指定关系）\ng(y)=(1−y/hc)^m；K=cosθ+r sinθ\nw₀=q/K；v₀=rw₀；w(y)=w₀g(y)，v(y)=v₀g(y)\ncrack_deformation｜未实现S式7的PST/Green求解',215,145,color=ORANGE),
N('vc','② 混凝土抗剪贡献（经验闭合，不是所给论文原式）\nVc=Vc0(1−exp(−q/qr))\n×[ρc+(1−ρc)exp(−w₀/wc)](1−Dc)\ntrial｜各经验参数需标定',405,140,color=ORANGE),
N('vs','③ 穿缝箍筋贡献\nεs,j=qg(yj)/ls；σs,j=min(Esεs,j,fy)；Vs=ΣAs,jσs,j\n应变传力长度ls为代码假定；损伤达到1时该筋应力归零\ntrial｜应力上限/离散求和有R22式9a、9c背景',590,140,color=GRAY),
N('bar','④ 对每根未失效FRP筋，计算裂缝运动投影\nδj=qg(yj)[sin(θ+αj)−r cos(θ+αj)]/K\n竖向筋αj=90°时：δj=qg(yj)\ntrial｜不是直接照搬S式8',775,140,color=GRAY),
N('tm','⑤ 由节点局部龄期更新 τm,i [F式2]\nτm,i=max{0,τm,0[1−kf log₁₀max(Ni,1)]}\nbond_strength；Ni为激活后局部循环数',960,115,color=GREEN),
N('bond','⑥ 三段黏结包络 [F式1]，令 tr=min(τr,τm,i)\nτenv=τm,i·s/δm                       s≤δm\nτenv=tr+(τm,i−tr)(δu−s)/(δu−δm)   δm<s≤δu\nτenv=tr                            s>δu\nbond：tr取min及历史割线规则为代码补充',1120,180,color=GREEN),
N('anchor','⑦ 双侧锚固离散求解与节点应力积分\ndT/dx=−pfτ；ds/dx=−T/(EfAf)；自由端T=0\nTi(x)=pf∫[x,Li]τi(ξ)dξ；anchorage / anchor_mesh\n积分原理：F式10、S式10；积分区间按本锚固段改写',1345,150,color=GREEN),
N('compat','⑧ 满足两侧协调与交点应变连续\ns₁(0)+s₂(0)=δj [S式9d]；εf1(0)=εf2(0) [S式11]\n同一根筋EfAf相同，因此T₁(0)=T₂(0)\n数值迭代求解，未直接使用F式9b或S式9a–c解析解',1540,150,color=GREEN),
N('fail','⑨ 检查本根筋脱黏、抗拉与用户保护层限值\n失效筋贡献归零；重复④—⑨处理其他筋\ntrial：筋轴力上限为Af·ffu，目前没有FRP本体疲劳律',1735,125,color=GRAY),
N('sum','⑩ 汇总竖向FRP抗剪力及整体平衡\nVf=nfaces ΣTj sinαj；V=Vc+Vs+Vf=Q/2\n总分解与E式1同一平衡思想；此处第三项是NSM筋\n不使用ECC覆盖层的 Vef,σ+Vef,τ 分解',1905,150,color=GREEN),
N('ok','局部求解有效，且\n|V(q)−Q/2|≤forceTolerance？',2100,115,'diamond',ORANGE),
N('return','返回该峰/谷的平衡解：q、材料分量与局部场\n供总控制器判断是否接受循环块',2265,100,'round',BLUE),
N('more','还有可继续的\n扫描/二分搜索？',2100,115,'diamond',ORANGE,x=1120,w=330),
N('diag','返回分类诊断\n无指定分支根 / 数值失败 / 分辨率限制\n未把“无解”直接归为抗弯IC脱黏破坏',2265,140,color=RED,x=1120,w=330)
]
le=[E('start','crack'),E('crack','vc'),E('vc','vs'),E('vs','bar'),E('bar','tm'),E('tm','bond'),E('bond','anchor'),E('anchor','compat'),E('compat','fail'),E('fail','sum'),E('sum','ok'),E('ok','return','是'),E('ok','more','否','r','l'),E('more','start','是：调整q','r','r',[[1510,2157],[1510,117]]),E('more','diag','否')]
for edge in le:
 if edge['a']=='more' and edge['b']=='start':edge['lp']=[1420,1000]
compare=[
N('a','你找出的公式，在当前抗剪代码中如何取舍',80,85,'round',BLUE,x=240,w=1080),
N('b','采用：F式(1)三段包络、F式(2)局部疲劳退化\n对应bond.m、bond_strength.m；0.75δm阈值用于active_nodes.m\n注意代码对零龄期、负峰值与残余值另有数值处理',230,150,color=GREEN,x=240,w=1080),
N('c','原理改写：F式(10)与S式(10)的黏结应力积分\n原文半跨坐标改为每根筋裂缝交点两侧的局部坐标：Ti(x)=pf∫[x,Li]τi dξ\nS式(9d)、(11)提供双侧滑移协调、交点应变连续的背景',460,170,color=GREEN,x=240,w=1080),
N('d','当前未直接实现：F式(9b)解析系数矩阵、S式(9a–c)分段解析解\nanchorage.m实际组装杆单元与黏结刚度，并进行固定点/Newton数值迭代\n两种矩阵写法不能仅因外观相似就标为同一个方程',710,170,color=ORANGE,x=240,w=1080),
N('e','当前不纳入：F式(11)的Ps/Pd、F式(13)的半跨Ls/Ld关系\n它们服务于纵向加固的抗弯IC脱黏模型；P也不是本抗剪算例的总加载力Q\n当前代码没有按这些式子计算临界荷载和半跨软化/脱黏长度',960,170,color=ORANGE,x=240,w=1080),
N('f','当前不纳入：F式(14)的纵向加固梁挠度\n竖向NSM抗剪不能直接使用该偏心纵向筋的弯曲挠度表达式\n当前的q、裂缝法向开口w与该文挠度w不是同一个物理量',1210,170,color=ORANGE,x=240,w=1080),
N('g','平衡思想可采用：E式(1) V=Vc+Vs+Vef\n在本模型改写为V=Vc+Vs+Vf；Vf来自实际NSM筋竖向分量\nE式(2)的Vef,σ+Vef,τ用于ECC连续覆盖层，当前NSM筋代码未实现',1460,180,color=GREEN,x=240,w=1080),
N('h','图示的是当前代码的计算路线，源码没有改变\n若要真正实现F式(9b)、(11)、(13)或S完整裂缝理论，需要另行建模与验证\n本轮不将这些尚未实现的公式画成已经运行的计算步骤',1730,160,color=BLUE,x=240,w=1080)
]
diagrams=[dict(name='公式对应版｜抗剪疲劳主循环',width=1560,height=2600,nodes=main,edges=edges),dict(name='公式对应版｜抗剪平衡内部求解',width=1560,height=2550,nodes=local,edges=le),dict(name='公式采用范围与当前实现',width=1560,height=2100,nodes=compare,edges=[])]
footers=['F＝2019抗弯疲劳modeling论文；S＝2019单裂缝抗剪论文；E＝2024 ECC-FRP抗剪论文。公式号必须带论文前缀。','F：DOI 10.1016/j.ijfatigue.2019.04.039；S：DOI尾号0000953；E：DOI尾号CCENG-4395。橙色包含附加假定。','R22：DOI尾号0001251，支持穿缝筋应力与离散贡献背景；本轮保留当前源码，仅重画公式对应的计算逻辑。']
(r/'diagram_data.json').write_text(json.dumps(diagrams,ensure_ascii=False,indent=2))
mx=Element('mxfile',host='app.diagrams.net',type='device')
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(28)
for index,d in enumerate(diagrams):
 nodes={n['id']:n for n in d['nodes']}
 def port(n,p):
  return {'t':(n['x']+n['w']/2,n['y']),'b':(n['x']+n['w']/2,n['y']+n['h']),'l':(n['x'],n['y']+n['h']/2),'r':(n['x']+n['w'],n['y']+n['h']/2)}[p]
 md=SubElement(mx,'diagram',name=d['name'],id=f'calc{index}')
 model=SubElement(md,'mxGraphModel',grid='1',gridSize='10',page='1',pageWidth=str(d['width']),pageHeight=str(d['height']))
 root=SubElement(model,'root');SubElement(root,'mxCell',id='0');SubElement(root,'mxCell',id='1',parent='0')
 def celltext(id,s,x,y,w,h,font=22):
  c=SubElement(root,'mxCell',id=id,value=s,style=f'text;html=0;align=center;verticalAlign=middle;whiteSpace=wrap;fontFamily=Microsoft YaHei;fontSize={font};',vertex='1',parent='1');SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
 celltext('title',d['name'],0,5,d['width'],45,28)
 celltext('footer',footers[index],25,d['height']-80,d['width']-50,65,18)
 for n in d['nodes']:
  shape='rhombus' if n['kind']=='diamond' else 'rectangle'
  style=f'shape={shape};rounded=1;whiteSpace=wrap;html=0;fillColor=#{n["color"]};strokeColor=#6A7E90;fontFamily=Microsoft YaHei;fontSize=21;'
  c=SubElement(root,'mxCell',id=n['id'],value=n['text'],style=style,vertex='1',parent='1');SubElement(c,'mxGeometry',x=str(n['x']),y=str(n['y']),width=str(n['w']),height=str(n['h']),attrib={'as':'geometry'})
 for i,e in enumerate(d['edges']):
  style='edgeStyle=orthogonalEdgeStyle;rounded=0;endArrow=block;endFill=1;strokeColor=#1685B6;fontColor=#B8302E;fontSize=19;'
  pp={'t':(.5,0),'b':(.5,1),'l':(0,.5),'r':(1,.5)};a=pp[e['ap']];b=pp[e['bp']];style+=f'exitX={a[0]};exitY={a[1]};entryX={b[0]};entryY={b[1]};'
  c=SubElement(root,'mxCell',id='e'+str(i),value=e['label'],style=style,edge='1',parent='1',source=e['a'],target=e['b']);geo=SubElement(c,'mxGeometry',relative='1',attrib={'as':'geometry'})
  if e['via']:
   arr=SubElement(geo,'Array',attrib={'as':'points'})
   for x,y in e['via']:SubElement(arr,'mxPoint',x=str(x),y=str(y))
 # Preview SVG uses identical editable vector primitives and labels.
 svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{d["width"]}" height="{d["height"]}" viewBox="0 0 {d["width"]} {d["height"]}">','<defs><marker id="a" markerWidth="9" markerHeight="9" refX="8" refY="4.5" orient="auto"><path d="M0,0 L9,4.5 L0,9Z" fill="#1685B6"/></marker></defs>','<rect width="100%" height="100%" fill="white"/>']
 def svgtext(s,cx,cy,fs=21,color='#243746'):
  lines=s.split('\n');out=[]
  for i,line in enumerate(lines):out.append(f'<text x="{cx}" y="{cy+(i-(len(lines)-1)/2)*fs*1.32}" font-family="Noto Sans CJK SC,Microsoft YaHei,sans-serif" font-size="{fs}" text-anchor="middle" dominant-baseline="middle" fill="{color}">{html.escape(line)}</text>')
  return out
 svg+=svgtext(d['name'],780,30,28)
 for e in d['edges']:
  pts=[port(nodes[e['a']],e['ap'])]+e['via']+[port(nodes[e['b']],e['bp'])]
  svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in pts)+'" fill="none" stroke="#1685B6" stroke-width="3" marker-end="url(#a)"/>')
  tip=pts[-1];prev=pts[-2];dx=tip[0]-prev[0];dy=tip[1]-prev[1];length=math.hypot(dx,dy)
  if length:
   ux,uy=dx/length,dy/length;bx,by=tip[0]-18*ux,tip[1]-18*uy
   svg.append(f'<polygon points="{tip[0]},{tip[1]} {bx-7*uy},{by+7*ux} {bx+7*uy},{by-7*ux}" fill="#1685B6"/>')
  if e['label']:
   p0,p1=pts[0],pts[1];lx,ly=e.get('lp',[(p0[0]+p1[0])/2+18,(p0[1]+p1[1])/2-12]);svg+=svgtext(e['label'],lx,ly,19,'#B8302E')
 for n in d['nodes']:
  x,y,w,h=n['x'],n['y'],n['w'],n['h']
  if n['kind']=='diamond':svg.append(f'<polygon points="{x+w/2},{y} {x+w},{y+h/2} {x+w/2},{y+h} {x},{y+h/2}" fill="#{n["color"]}" stroke="#6A7E90"/>')
  else:svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="14" fill="#{n["color"]}" stroke="#6A7E90"/>')
  svg+=svgtext(n['text'],x+w/2,y+h/2,21 if n['w']>400 else 19)
 svg+=svgtext(footers[index],780,d['height']-45,17);svg.append('</svg>');(r/(d['name']+'.svg')).write_text('\n'.join(svg))
 # Standard PowerPoint nodes and native straight-segment connectors; no custom paths.
 slide=prs.slides.add_slide(prs.slide_layouts[6]);scale=min(15/d['width'],26.4/d['height']);ox=(16-d['width']*scale)/2;oy=(28-d['height']*scale)/2
 def pos(x,y):return Inches(ox+x*scale),Inches(oy+y*scale)
 def ppttext(sh,s,fs=21):
  tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.margin_top=tf.margin_bottom=0;tf.margin_left=tf.margin_right=0
  for i,t in enumerate(s.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t;p.alignment=PP_ALIGN.CENTER;p.font.name='Microsoft YaHei';p.font.size=Pt(fs*scale*72);p.space_after=Pt(1);p.font.color.rgb=RGBColor.from_string('243746')
 shapes={}
 for n in d['nodes']:
  kind=MSO_SHAPE.DIAMOND if n['kind']=='diamond' else MSO_SHAPE.ROUNDED_RECTANGLE
  sh=slide.shapes.add_shape(kind,*pos(n['x'],n['y']),Inches(n['w']*scale),Inches(n['h']*scale));sh.name=n['id'];sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(n['color']);sh.line.color.rgb=RGBColor.from_string('6A7E90');ppttext(sh,n['text'],21 if n['w']>400 else 19);shapes[n['id']]=sh
 for e in d['edges']:
  pts=[port(nodes[e['a']],e['ap'])]+e['via']+[port(nodes[e['b']],e['bp'])]
  for seg,(p0,p1) in enumerate(zip(pts,pts[1:])):
   line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*pos(*p0),*pos(*p1));line.name=e['a']+' → '+e['b']+f' / {seg+1}';line.line.color.rgb=RGBColor.from_string('1685B6');line.line.width=Pt(1.4)
   sites={'t':0,'l':1,'b':2,'r':3}
   if seg==0:line.begin_connect(shapes[e['a']],sites[e['ap']])
   if seg==len(pts)-2:
    line.end_connect(shapes[e['b']],sites[e['bp']]);tip=OxmlElement('a:tailEnd');tip.set('type','triangle');line._element.spPr.get_or_add_ln().append(tip)
   tree=slide.shapes._spTree;tree.remove(line._element);tree.insert(2,line._element)
  if e['label']:
   p0,p1=pts[0],pts[1];x,y=e.get('lp',[(p0[0]+p1[0])/2+18,(p0[1]+p1[1])/2-20])
   tb=slide.shapes.add_textbox(*pos(x-100,y-15),Inches(200*scale),Inches(35*scale));ppttext(tb,e['label'],18)
 for txt,x,y,w,h,fs in [(d['name'],0,5,1560,45,28),(footers[index],10,d['height']-80,1540,65,17)]:
  tb=slide.shapes.add_textbox(*pos(x,y),Inches(w*scale),Inches(h*scale));ppttext(tb,txt,fs)
ElementTree(mx).write(r/'公式对应版_抗剪疲劳计算逻辑_可编辑.drawio',encoding='utf-8',xml_declaration=True)
prs.save(r/'公式对应版_抗剪疲劳计算逻辑_可编辑.pptx')
print('Built three formula-mapped diagrams, PPTX and drawio.')
