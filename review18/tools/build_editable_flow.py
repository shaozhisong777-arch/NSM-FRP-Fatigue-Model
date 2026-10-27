"""Rebuild both DOT diagrams as native PowerPoint shapes and connected vector paths."""
from pathlib import Path
import json, re, subprocess
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
prs = Presentation()
# A large portrait canvas keeps the complete logic editable on two pages.
prs.slide_width = Inches(16)
prs.slide_height = Inches(24)
EMU = 914400

def element(name, **attrs):
    e = OxmlElement(name)
    for key, value in attrs.items():
        e.set(key, str(value))
    return e

def closest_site(shape, pt):
    x,y = pt
    points = [(shape.left+shape.width/2,shape.top),
              (shape.left,shape.top+shape.height/2),
              (shape.left+shape.width/2,shape.top+shape.height),
              (shape.left+shape.width,shape.top+shape.height/2)]
    return min(range(4),key=lambda k:(points[k][0]-x)**2+(points[k][1]-y)**2)

for filename in ['抗剪疲劳实现逻辑','局部力学计算']:
    raw = subprocess.check_output(['dot','-Tjson',str(ROOT/'figures'/(filename+'.dot'))])
    g = json.loads(raw)
    x0,y0,x1,y1 = map(float,g['bb'].split(','))
    width,height=x1-x0,y1-y0
    margin=.55
    scale=min((16-2*margin)/width,(24-2*margin)/height)
    left=(16-width*scale)/2
    top=(24-height*scale)/2
    def xy(point):
        return (round((left+(point[0]-x0)*scale)*EMU),
                round((top+(y1-point[1])*scale)*EMU))
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    # Draw native connectors behind the nodes, retaining the original Bezier routes.
    links=[]
    objects={o['_gvid']:o for o in g['objects'] if 'pos' in o}
    for edge in g['edges']:
        paths=[o for o in edge.get('_draw_',[]) if o['op'] in ['b','B']]
        if not paths:
            continue
        points=[xy(p) for p in paths[0]['points']]
        # Graphviz stops the spline before its arrowhead; extend it to the tip.
        pos=edge.get('pos','')
        tip=None
        if pos.startswith('e,'):
            v=pos.split()[0].split(',')
            tip=xy([float(v[1]),float(v[2])])
        allpoints=points+([tip] if tip else [])
        xmin=min(p[0] for p in allpoints);ymin=min(p[1] for p in allpoints)
        xmax=max(p[0] for p in allpoints);ymax=max(p[1] for p in allpoints)
        w=max(1,xmax-xmin);h=max(1,ymax-ymin)
        line=slide.shapes.add_connector(MSO_CONNECTOR.CURVE,xmin,ymin,xmin+w,ymin+h)
        line.name='流程箭头 '+objects[edge['tail']]['name']+' → '+objects[edge['head']]['name']
        spPr=line._element.spPr
        for child in list(spPr):
            if child.tag.endswith('prstGeom') or child.tag.endswith('custGeom'):
                spPr.remove(child)
        geom=element('a:custGeom')
        for tag in ['a:avLst','a:gdLst','a:ahLst','a:cxnLst']:
            geom.append(element(tag))
        geom.append(element('a:rect',l='0',t='0',r='r',b='b'))
        pathlist=element('a:pathLst');path=element('a:path',w=w,h=h,fill='none',stroke='1')
        def pt(p):
            return element('a:pt',x=p[0]-xmin,y=p[1]-ymin)
        move=element('a:moveTo');move.append(pt(points[0]));path.append(move)
        for i in range(1,len(points),3):
            if len(points[i:i+3])!=3:break
            bezier=element('a:cubicBezTo')
            for p in points[i:i+3]:bezier.append(pt(p))
            path.append(bezier)
        if tip:
            end=element('a:lnTo');end.append(pt(tip));path.append(end)
        pathlist.append(path);geom.append(pathlist);spPr.insert(1,geom)
        line.line.color.rgb=RGBColor.from_string('60758A');line.line.width=Pt(1.2)
        ln=spPr.get_or_add_ln()
        ln.append(element('a:tailEnd',type='triangle',w='sm',len='sm'))
        links.append((line,edge,points[0],tip or points[-1]))
    # Nodes are actual rounded rectangles/diamonds with editable text.
    nodes={}
    for gid,o in objects.items():
        cx,cy=map(float,o['pos'].split(','))
        nw=float(o['width'])*72;nh=float(o['height'])*72
        x,y=xy([cx-nw/2,cy+nh/2])
        w=round(nw*scale*EMU);h=round(nh*scale*EMU)
        kind=MSO_SHAPE.DIAMOND if o.get('shape')=='diamond' else MSO_SHAPE.ROUNDED_RECTANGLE
        shape=slide.shapes.add_shape(kind,x,y,w,h)
        shape.name='节点 '+o['name']+' | '+o.get('label','').split('\\')[0]
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(o.get('fillcolor','#eaf2ff').lstrip('#'))
        shape.line.color.rgb=RGBColor.from_string(o.get('color','#4b6f99').lstrip('#'));shape.line.width=Pt(1)
        tf=shape.text_frame;tf.clear();tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        tf.margin_top=tf.margin_bottom=0;tf.margin_left=tf.margin_right=Inches(.04)
        lines=[v['text'] for v in o.get('_ldraw_',[]) if v['op']=='T']
        fs=float(o.get('fontsize',12))*scale*72
        for i,label in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=label;p.alignment=PP_ALIGN.CENTER
            p.font.name='Microsoft YaHei';p.font.size=Pt(fs)
            p.font.color.rgb=RGBColor.from_string('182F45')
            p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(fs*1.12)
        nodes[gid]=shape
    for line,edge,start,end in links:
        transform = deepcopy(line._element.spPr.xfrm)
        line.begin_connect(nodes[edge['tail']],closest_site(nodes[edge['tail']],start))
        line.end_connect(nodes[edge['head']],closest_site(nodes[edge['head']],end))
        # Keep the Graphviz route while recording native endpoint associations.
        current = line._element.spPr.xfrm
        current.getparent().replace(current, transform)
    # Edge labels and diagram headings remain independent editable text boxes.
    for obj in [g]+g['edges']:
        size=12
        for op in obj.get('_ldraw_',[]):
            if op['op']=='F':size=op['size']
            if op['op']!='T':continue
            tw=max(op.get('width',40)+12,24)
            th=size*1.55
            x,y=xy([op['pt'][0]-tw/2,op['pt'][1]+size*1.05])
            box=slide.shapes.add_textbox(x,y,round(tw*scale*EMU),round(th*scale*EMU))
            box.name='可编辑标签 | '+op['text']
            tf=box.text_frame;tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            p=tf.paragraphs[0];p.text=op['text'];p.alignment=PP_ALIGN.CENTER
            p.font.name='Microsoft YaHei';p.font.size=Pt(size*scale*72)
            p.font.color.rgb=RGBColor.from_string('243746')
    # Notes tell the user that this is editable, not an embedded screenshot.
    notes=slide.notes_slide.notes_text_frame
    notes.text=('所有节点、文字、公式和箭头均为PowerPoint原生对象，没有插入逻辑图图片。'
                '在“开始→选择→选择窗格”按节点或箭头名称选择对象。'
                '箭头为带节点连接关系的自定义矢量连接符；移动节点后请检查回路走向，可重新路由或编辑顶点。'
                '画布16×24英寸，建议放大后编辑；公式为可编辑Unicode文本。'
                'F编号与review18中文审查报告一致。')

output=ROOT/'抗剪疲劳逻辑图_完全可编辑_CN.pptx'
prs.save(output)
print(output)
